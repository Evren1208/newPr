from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .analysis_service import class_overview, compare_exams, subject_stats
from .models import ExamData


def export_analysis(current: ExamData, previous: ExamData | None, out_path: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "班级概览"
    overview = class_overview(current)
    rows = [
        ("考试", current.exam_name),
        ("学生人数", int(overview["student_count"])),
        ("科目数", int(overview["subject_count"])),
        ("人均总分", round(overview["avg_total"], 2)),
        ("总分中位数", round(overview["median_total"], 2)),
        ("最高总分", round(overview["max_total"], 2)),
        ("最低总分", round(overview["min_total"], 2)),
        ("人均单科分", round(overview["avg_subject_score"], 2)),
    ]
    for i, (k, v) in enumerate(rows, start=1):
        ws.cell(i, 1, k)
        ws.cell(i, 2, v)
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 22

    ws2 = wb.create_sheet("学科统计")
    headers = ["科目", "平均分", "中位数", "最高分", "最低分", "及格率%", "优秀率%"]
    ws2.append(headers)
    for cell in ws2[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.fill = PatternFill("solid", fgColor="EDEDED")
    for s in subject_stats(current):
        ws2.append([s.subject, round(s.mean, 2), round(s.median, 2), s.max_score, s.min_score,
                    round(s.pass_rate, 2), round(s.excellent_rate, 2)])
    for col in "ABCDEFG":
        ws2.column_dimensions[col].width = 16

    ws3 = wb.create_sheet("学生成绩")
    ws3.append(["姓名"] + current.subjects + ["总分", "平均分"])
    for cell in ws3[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EDEDED")
    for st in current.students:
        ws3.append([st.name] + [st.scores.get(sub, "") for sub in current.subjects] + [st.total, round(st.average, 2)])

    changes = compare_exams(current, previous)
    if changes:
        ws4 = wb.create_sheet("进退步")
        ws4.append(["姓名", "本次总分", "上次总分", "变化"])
        for c in changes:
            ws4.append([c.name, c.current_total, c.previous_total, c.delta])
        for col in "ABCD":
            ws4.column_dimensions[col].width = 18

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return str(out)
