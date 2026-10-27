from __future__ import annotations

import re
from pathlib import Path
from typing import Dict, List, Tuple

from openpyxl import load_workbook

from .models import ExamData, MappingConfig, StudentRecord

NAME_ALIASES = {"姓名", "学生姓名", "名字", "name", "student name"}
ID_ALIASES = {"学号", "学生编号", "编号", "id", "student id"}
CLASS_ALIASES = {"班级", "行政班", "class"}
EXCLUDE_NUMERIC = {
    "总分", "总成绩", "平均分", "均分", "排名", "名次", "班级排名", "年级排名",
    "序号", "座号", "学号", "学生编号", "编号", "总评"
}
COMMON_SUBJECTS = {
    "语文", "数学", "英语", "物理", "化学", "生物", "政治", "思想政治", "历史", "地理",
    "科学", "道法", "道德与法治", "信息技术", "体育", "音乐", "美术"
}


def _norm(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    s = re.sub(r"\s+", "", s)
    return s.lower()


def _find_header_row(ws, max_scan: int = 15) -> int:
    best_row = 1
    best_score = -1
    for r in range(1, min(max_scan, ws.max_row) + 1):
        vals = [_norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        nonempty = [v for v in vals if v]
        score = 0
        if any(v in NAME_ALIASES for v in nonempty):
            score += 8
        score += min(len(nonempty), 12)
        score += sum(2 for v in nonempty if v in COMMON_SUBJECTS)
        if score > best_score:
            best_score = score
            best_row = r
    return best_row


def _to_float(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    s = str(v).strip().replace(",", "").replace("分", "")
    try:
        return float(s)
    except ValueError:
        return None


def _numeric_ratio(ws, col: int, start_row: int, sample_size: int = 60) -> float:
    total = 0
    numeric = 0
    for r in range(start_row, min(ws.max_row, start_row + sample_size - 1) + 1):
        v = ws.cell(r, col).value
        if v is None or str(v).strip() == "":
            continue
        total += 1
        if _to_float(v) is not None:
            numeric += 1
    return numeric / total if total else 0.0


def inspect_excel(path: str) -> Dict:
    wb = load_workbook(path, data_only=True, read_only=False)
    ws = wb[wb.sheetnames[0]]
    header_row = _find_header_row(ws)
    headers: List[str] = []
    normalized: Dict[int, str] = {}
    for c in range(1, ws.max_column + 1):
        raw = ws.cell(header_row, c).value
        label = str(raw).strip() if raw is not None else f"列{c}"
        headers.append(label)
        normalized[c] = _norm(label)

    name_col = next((c for c, v in normalized.items() if v in NAME_ALIASES), None)
    if name_col is None:
        for c in range(1, ws.max_column + 1):
            if _numeric_ratio(ws, c, header_row + 1) < 0.35:
                name_col = c
                break

    id_col = next((c for c, v in normalized.items() if v in ID_ALIASES), None)
    class_col = next((c for c, v in normalized.items() if v in CLASS_ALIASES), None)

    subject_cols: List[int] = []
    exclude_norm = {_norm(x) for x in EXCLUDE_NUMERIC}
    subject_norm = {_norm(x) for x in COMMON_SUBJECTS}
    for c in range(1, ws.max_column + 1):
        if c in {name_col, id_col, class_col}:
            continue
        label = headers[c - 1]
        n = normalized[c]
        if n in exclude_norm:
            continue
        ratio = _numeric_ratio(ws, c, header_row + 1)
        if n in subject_norm or ratio >= 0.65:
            if any(k in label for k in ["排名", "名次", "序号", "座号", "学号", "总分", "平均"]):
                continue
            subject_cols.append(c)

    return {
        "sheet": ws.title,
        "header_row": header_row,
        "headers": headers,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "suggested": MappingConfig(
            header_row=header_row,
            name_col=name_col or 1,
            subject_cols=subject_cols,
            id_col=id_col,
            class_col=class_col,
        ),
    }


def _read_with_mapping(path: str, exam_name: str, mapping: MappingConfig) -> ExamData:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(path)

    wb = load_workbook(path, data_only=True, read_only=False)
    ws = wb[wb.sheetnames[0]]
    header_row = mapping.header_row

    headers: Dict[int, str] = {}
    for c in range(1, ws.max_column + 1):
        raw = ws.cell(header_row, c).value
        headers[c] = str(raw).strip() if raw is not None else f"列{c}"

    if not mapping.subject_cols:
        raise ValueError("请至少选择一个成绩科目列。")

    subject_cols: List[Tuple[int, str]] = [(c, headers[c]) for c in mapping.subject_cols]
    students: List[StudentRecord] = []
    warnings: List[str] = []

    for r in range(header_row + 1, ws.max_row + 1):
        name_raw = ws.cell(r, mapping.name_col).value
        if name_raw is None or str(name_raw).strip() == "":
            continue
        name = str(name_raw).strip()
        scores: Dict[str, float] = {}
        for c, subject in subject_cols:
            value = _to_float(ws.cell(r, c).value)
            if value is not None:
                scores[subject] = value
        if not scores:
            continue
        sid = None
        class_name = None
        if mapping.id_col:
            value = ws.cell(r, mapping.id_col).value
            sid = str(value).strip() if value is not None else None
        if mapping.class_col:
            value = ws.cell(r, mapping.class_col).value
            class_name = str(value).strip() if value is not None else None
        students.append(StudentRecord(name=name, scores=scores, student_id=sid, class_name=class_name))

    if len(students) < 2:
        warnings.append("识别到的学生数量较少，请检查字段映射是否正确。")

    return ExamData(
        students=students,
        subjects=[label for _, label in subject_cols],
        source_file=str(p),
        exam_name=exam_name,
        header_row=header_row,
        warnings=warnings,
    )


def read_exam(path: str, exam_name: str = "本次考试", mapping: MappingConfig | None = None) -> ExamData:
    info = inspect_excel(path)
    use_mapping = mapping or info["suggested"]
    return _read_with_mapping(path, exam_name, use_mapping)
