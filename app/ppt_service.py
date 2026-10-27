from __future__ import annotations

from pathlib import Path
from typing import Optional

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from .analysis_service import class_overview, compare_exams, rule_based_summary, score_distribution, subject_stats
from .models import ExamData

FONT = "Microsoft YaHei"


def _set_font(run, size=20, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _add_title(slide, title: str, subtitle: str | None = None):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(12), Inches(0.7))
    p = box.text_frame.paragraphs[0]
    p.text = title
    _set_font(p.runs[0], 28, True)
    if subtitle:
        box2 = slide.shapes.add_textbox(Inches(0.72), Inches(1.15), Inches(11.5), Inches(0.45))
        p2 = box2.text_frame.paragraphs[0]
        p2.text = subtitle
        _set_font(p2.runs[0], 13, False)


def _add_bullets(slide, items, top=1.6):
    box = slide.shapes.add_textbox(Inches(0.9), Inches(top), Inches(11.5), Inches(5.2))
    tf = box.text_frame
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.space_after = Pt(12)
        _set_font(p.runs[0], 20 if idx == 0 else 18, idx == 0)


def generate_ppt(current: ExamData, previous: Optional[ExamData], out_path: str,
                 class_title: str = "班级家长会", ai_summary: str | None = None) -> str:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    s = _blank(prs)
    title = s.shapes.add_textbox(Inches(1.0), Inches(2.1), Inches(11.3), Inches(1.2))
    p = title.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = class_title
    _set_font(p.runs[0], 34, True)
    sub = s.shapes.add_textbox(Inches(1.0), Inches(3.45), Inches(11.3), Inches(0.8))
    p2 = sub.text_frame.paragraphs[0]
    p2.alignment = PP_ALIGN.CENTER
    p2.text = f"{current.exam_name} · 自动生成"
    _set_font(p2.runs[0], 18, False)

    ov = class_overview(current)
    s = _blank(prs)
    _add_title(s, "01  本次考试概览")
    _add_bullets(s, [
        f"参加分析学生：{int(ov['student_count'])} 人",
        f"纳入统计科目：{int(ov['subject_count'])} 科",
        f"人均总分：{ov['avg_total']:.1f} 分",
        f"最高总分：{ov['max_total']:.1f} 分；最低总分：{ov['min_total']:.1f} 分",
        f"人均单科成绩：{ov['avg_subject_score']:.1f} 分",
    ])

    stats = subject_stats(current)
    s = _blank(prs)
    _add_title(s, "02  各学科平均分")
    cd = ChartData()
    cd.categories = [x.subject for x in stats]
    cd.add_series("平均分", [round(x.mean, 1) for x in stats])
    chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.9), Inches(1.6), Inches(11.6), Inches(4.8), cd).chart
    chart.has_legend = False
    chart.value_axis.maximum_scale = 100
    chart.value_axis.minimum_scale = 0

    s = _blank(prs)
    _add_title(s, "03  及格率与优秀率")
    cd = ChartData()
    cd.categories = [x.subject for x in stats]
    cd.add_series("及格率", [round(x.pass_rate, 1) for x in stats])
    cd.add_series("优秀率", [round(x.excellent_rate, 1) for x in stats])
    chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.9), Inches(1.6), Inches(11.6), Inches(4.8), cd).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.value_axis.maximum_scale = 100
    chart.value_axis.minimum_scale = 0

    dist = score_distribution(current)
    s = _blank(prs)
    _add_title(s, "04  学生平均分分布")
    cd = ChartData()
    cd.categories = list(dist.keys())
    cd.add_series("人数", list(dist.values()))
    chart = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.9), Inches(1.6), Inches(11.6), Inches(4.8), cd).chart
    chart.has_legend = False

    s = _blank(prs)
    _add_title(s, "05  学习变化")
    changes = compare_exams(current, previous)
    if changes:
        top_up = [x for x in changes if x.delta > 0][:5]
        top_down = [x for x in sorted(changes, key=lambda x: x.delta) if x.delta < 0][:5]
        items = []
        if top_up:
            items.append("班主任内部参考 · 进步较明显：" + "、".join(f"{x.name}（{x.delta:+.0f}）" for x in top_up))
        if top_down:
            items.append("班主任内部参考 · 需要关注：" + "、".join(f"{x.name}（{x.delta:+.0f}）" for x in top_down))
        items.append("公开家长会版本建议删除学生姓名与个人排名，仅展示班级总体统计。")
    else:
        items = ["当前未导入上次考试成绩。", "导入上一份成绩表后，可自动生成学生进退步与班级趋势。"]
    _add_bullets(s, items)

    summary = rule_based_summary(current, previous)
    s = _blank(prs)
    _add_title(s, "06  班级表现总结")
    _add_bullets(s, [ai_summary] if ai_summary else [summary["highlights"], summary["concerns"]])

    s = _blank(prs)
    _add_title(s, "07  下一阶段重点")
    _add_bullets(s, [summary["suggestions"], "建议按学生基础分层设置短周期目标，并在下一次考试后复盘变化。"])

    s = _blank(prs)
    _add_title(s, "08  家庭配合建议")
    _add_bullets(s, [
        "关注学习节奏，而不仅是一次考试分数。",
        "帮助孩子形成固定的作息、复习与错题整理习惯。",
        "遇到明显波动时，先寻找原因，再和孩子共同确定改进目标。",
        "家校保持必要沟通，避免公开比较学生个人成绩和排名。",
    ])

    s = _blank(prs)
    title = s.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(11.3), Inches(1.4))
    p = title.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "感谢各位家长的理解与配合"
    _set_font(p.runs[0], 30, True)

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return str(out)
