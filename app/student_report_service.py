from __future__ import annotations

from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from .analysis_service import student_comment
from .models import ExamData


def _set_run_font(run, size=11, bold=False):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold


def generate_student_report(current: ExamData, previous: Optional[ExamData], out_path: str) -> str:
    doc = Document()
    prev_map = {s.name: s for s in previous.students} if previous else {}

    for idx, student in enumerate(current.students):
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run(f"{student.name} · 学习情况报告")
        _set_run_font(run, 16, True)

        p = doc.add_paragraph()
        run = p.add_run(f"本次总分：{student.total:.1f}    平均分：{student.average:.1f}")
        _set_run_font(run, 11, True)

        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "科目"
        hdr[1].text = "本次"
        hdr[2].text = "上次"
        hdr[3].text = "变化"

        prev = prev_map.get(student.name)
        for subject in current.subjects:
            row = table.add_row().cells
            current_score = student.scores.get(subject)
            previous_score = prev.scores.get(subject) if prev else None
            delta = current_score - previous_score if current_score is not None and previous_score is not None else None
            row[0].text = subject
            row[1].text = "" if current_score is None else f"{current_score:.1f}"
            row[2].text = "" if previous_score is None else f"{previous_score:.1f}"
            row[3].text = "" if delta is None else f"{delta:+.1f}"

        doc.add_paragraph()
        p = doc.add_paragraph()
        r = p.add_run("学习建议：")
        _set_run_font(r, 11, True)
        r = p.add_run(student_comment(student.name, student.scores, prev.scores if prev else None))
        _set_run_font(r, 11, False)

        note = doc.add_paragraph()
        r = note.add_run("说明：本报告用于学生本人及家长了解学习情况，不建议公开展示。")
        _set_run_font(r, 9, False)

        if idx < len(current.students) - 1:
            doc.add_page_break()

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return str(out)
